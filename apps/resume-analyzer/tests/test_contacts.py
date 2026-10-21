import json
import unittest
from io import BytesIO
from unittest.mock import patch

from docx import Document
from fastapi import HTTPException, UploadFile

from app.candidate_extractor import extract_contacts
from app.job_processor import process_job
from app.main import rank_resume, resume_contacts
from app.rank_models import RankResponse
from app.readiness import mark_ready


class ContactExtractionTests(unittest.TestCase):
    def test_wrapped_patronymic_has_exact_source(self):
        result = extract_contacts("Румянцева Дарья\nКонстантиновна\nТелефон: +7 (999) 123-45-67")
        self.assertEqual(result["candidate"], {"last_name": "Румянцева", "first_name": "Дарья",
                                              "middle_name": "Константиновна", "phone": "+79991234567"})
        source = result["contact_details"]["fields"]["middle_name"]
        self.assertEqual(source, {"source": "resume_text", "confidence": "high",
                                  "quote": "Румянцева Дарья\nКонстантиновна", "line_start": 1, "line_end": 2})

    def test_surname_on_separate_line_and_three_lines(self):
        for text in ("Румянцева\nДарья Константиновна", "Румянцева\nДарья\nКонстантиновна",
                     "РЕЗЮМЕ\nРумянцева\n\nДарья\nКонстантиновна", "Дарья Константиновна\nРумянцева"):
            with self.subTest(text=text):
                self.assertEqual(extract_contacts(text)["candidate"]["middle_name"], "Константиновна")

    def test_absent_patronymic_stays_null(self):
        result = extract_contacts("Румянцева Дарья\nОпыт работы\nВысшее образование")
        self.assertEqual(result["candidate"]["last_name"], "Румянцева")
        self.assertIsNone(result["candidate"]["middle_name"])
        self.assertIsNone(result["contact_details"]["fields"]["middle_name"])

    def test_uppercase_and_first_name_before_surname(self):
        for text in ("РУМЯНЦЕВА ДАРЬЯ КОНСТАНТИНОВНА", "Дарья Румянцева Константиновна"):
            self.assertEqual(extract_contacts(text)["candidate"]["first_name"], "Дарья")
            self.assertEqual(extract_contacts(text)["candidate"]["last_name"], "Румянцева")

    def test_headings_and_job_titles_are_not_names(self):
        for text in ("Опыт Работы\nВысшее Образование", "Ключевые Навыки\nЛичные Данные",
                     "Менеджер Проектов", "Старший Аналитик", "Русский Язык"):
            self.assertIsNone(extract_contacts(text)["candidate"]["first_name"], text)

    def test_unfamiliar_labelled_name_requires_review(self):
        result = extract_contacts("ФИО: Румянцева Дарина Константиновна")
        self.assertEqual(result["candidate"]["first_name"], "Дарина")
        self.assertEqual(result["contact_details"]["fields"]["first_name"]["confidence"], "medium")

    def test_label_and_each_name_part_on_separate_lines(self):
        result = extract_contacts("ФИО:\nРумянцева\nДарина\nКонстантиновна")
        self.assertEqual(result["candidate"]["middle_name"], "Константиновна")
        self.assertEqual(result["contact_details"]["fields"]["middle_name"]["line_end"], 4)
        self.assertEqual(result["contact_details"]["fields"]["first_name"]["confidence"], "medium")

    def test_multiple_names_are_ambiguous(self):
        result = extract_contacts("Румянцева Дарья\nИванов Иван")
        self.assertEqual(result["contact_details"]["name_status"], "ambiguous")
        self.assertIsNone(result["candidate"]["last_name"])

    def test_russian_phone_prefixes(self):
        for phone in ("+7 (999) 123-45-67", "8 999 123 45 67", "79991234567", "+7\u00a0999\u202f123‑45‑67"):
            result = extract_contacts("Телефон: " + phone)
            self.assertEqual(result["candidate"]["phone"], "+79991234567")
            self.assertEqual(result["contact_details"]["phone_status"], "extracted")

    def test_unsupported_international_short_and_local_numbers_are_explicit(self):
        for phone in ("+44 20 7946 0958", "+1 (212) 555-0123", "9991234567", "123-45-67", "+7 999"):
            with self.subTest(phone=phone):
                result = extract_contacts("Телефон: " + phone)
                self.assertIsNone(result["candidate"]["phone"])
                self.assertEqual(result["contact_details"]["phone_status"], "unsupported")
                self.assertTrue(result["contact_details"]["warnings"])

    def test_multiple_phones_do_not_silently_select_one(self):
        result = extract_contacts("Телефон: +7 999 123 45 67\nТелефон: +7 999 123 45 68")
        self.assertEqual(result["contact_details"]["phone_status"], "ambiguous")
        self.assertIsNone(result["candidate"]["phone"])

    def test_occupation_and_mobile_development_do_not_invalidate_phone(self):
        for experience in ("Преподаватель 2019 - 2026", "Мобильная разработка 2019 - 2026",
                           "Мобильный разработчик 2019 - 2026"):
            with self.subTest(experience=experience):
                result = extract_contacts("Телефон: +7 999 123 45 67\n" + experience)
                self.assertEqual(result["candidate"]["phone"], "+79991234567")
                self.assertEqual(result["contact_details"]["phone_status"], "extracted")

    def test_unsupported_phone_below_label_is_explicit(self):
        for text in ("Телефон:\n9991234567", "Моб.:\n\n123-45-67", "+7 999"):
            with self.subTest(text=text):
                result = extract_contacts(text)
                self.assertIsNone(result["candidate"]["phone"])
                self.assertEqual(result["contact_details"]["phone_status"], "unsupported")

    def test_dates_salary_and_missing_phone_are_not_numbers(self):
        result = extract_contacts("Румянцева Дарья\n2019 - 2026\nЗарплата: 100000\nДата рождения 01.01.1990")
        self.assertEqual(result["contact_details"]["phone_status"], "not_found")
        self.assertIsNone(result["candidate"]["phone"])

    def test_empty_text_is_explicit(self):
        result = extract_contacts(" \n\t")
        self.assertTrue(all(value is None for value in result["candidate"].values()))
        self.assertEqual(result["contact_details"]["text_status"], "no_text")
        self.assertTrue(result["contact_details"]["warnings"])


class ContactPathTests(unittest.TestCase):
    def setUp(self):
        mark_ready(True)

    def tearDown(self):
        mark_ready(False)

    def test_http_amqp_and_recovery_return_the_same_contacts(self):
        text = "Румянцева Дарья\nКонстантиновна\nТелефон: 8 (999) 123-45-67\nОпыт работы: аналитик данных."
        content = BytesIO()
        doc = Document()
        for line in text.splitlines():
            doc.add_paragraph(line)
        doc.save(content)
        data = content.getvalue()
        job = {"version": 1, "job_id": "job", "batch_id": "batch", "item_id": "item", "attempt": 1,
               "file": {"file_name": "resume.docx"}, "requirements": [{"id": "cluster", "name": "Test", "profiles": []}]}
        # The HTTP path requires at least one profile; no actual model calls are needed here.
        job["requirements"][0]["profiles"] = [{"id": "profile", "position_name": "Analyst"}]
        rating = RankResponse.model_validate({"clusters": [{"id": "cluster", "name": "Test", "profiles": [{
            "id": "profile", "position_name": "Analyst", "score": 0, "explanation": "Test"}]}]})
        upload = lambda: UploadFile(filename="resume.docx", file=BytesIO(data))
        with patch("app.main.is_ready", return_value=True), patch("app.main.analyze_resume_text", return_value=rating), \
             patch("app.job_processor.analyze_resume_text", return_value=rating), \
             patch("app.job_processor.download_document", return_value=data):
            http = rank_resume(upload(), json.dumps({"clusters": job["requirements"]})).model_dump()
            amqp = process_job(job)
        with patch("app.main.is_ready", return_value=False) as ready, patch("app.main.analyze_resume_text") as model:
            recovered = resume_contacts(upload()).model_dump()
            ready.assert_not_called()
            model.assert_not_called()
        for key in ("candidate", "contact_details"):
            self.assertEqual(http[key], amqp[key])
            self.assertEqual(http[key], recovered[key])

    def test_empty_contact_recovery_works_without_model_and_keeps_null(self):
        content = BytesIO()
        Document().save(content)
        with patch("app.main.is_ready", return_value=False), patch("app.main.analyze_resume_text") as model:
            result = resume_contacts(UploadFile(filename="empty.docx", file=BytesIO(content.getvalue()))).model_dump()
        self.assertIsNone(result["candidate"]["middle_name"])
        self.assertEqual(result["contact_details"]["text_status"], "no_text")
        model.assert_not_called()

    def test_invalid_document_is_a_controlled_422(self):
        with self.assertRaises(HTTPException) as caught:
            resume_contacts(UploadFile(filename="broken.docx", file=BytesIO(b"invalid zip")))
        self.assertEqual(caught.exception.status_code, 422)
