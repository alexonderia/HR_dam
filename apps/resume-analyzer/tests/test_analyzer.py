import json
import asyncio
import threading
import unittest
import time
from io import BytesIO
from unittest.mock import patch, Mock

import requests
from fastapi import HTTPException, UploadFile
from docx import Document

from app.main import app, lifespan, rank_resume, health_check, readiness
from app.ollama_client import (
    ask_ollama_chat,
    ModelUnavailable,
    ModelTimeout,
    ModelResponseError,
)
from app.rank_models import ProfileRequest
from app.scoring import (
    parse_ranking_response,
    rank_profile,
    build_ranking_prompt,
    ContextFitError,
)
from app.readiness import mark_ready, model_lock, monitor
from app.text_extractor import extract_text_from_file


class AnalyzerTests(unittest.TestCase):
    def setUp(self):
        mark_ready(True)
        self.profile = ProfileRequest(id="profile", position_name="Разработчик", keywords=["Python"])
        self.requirements = json.dumps({"clusters": [{"id": "cluster", "name": "ИТ",
            "comment": "Общие требования", "profiles": [self.profile.model_dump()]}]})

    def tearDown(self):
        mark_ready(False)

    def upload(self):
        doc = Document()
        doc.add_paragraph("Python SQL developer with five years of experience building REST APIs")
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        return UploadFile(filename="resume.docx", file=buf)

    def test_parse_rejects_invalid_scores_and_arrays(self):
        for score in ["80", True, -1, 101, float("nan"), float("inf")]:
            self.assertIsNone(parse_ranking_response(json.dumps({"score": score, "explanation": "OK"})))
        self.assertIsNone(parse_ranking_response('{"score":50,"explanation":"OK","matched_keywords":"Python"}'))
        self.assertEqual(parse_ranking_response('```json\n{"score":80,"explanation":"OK"}\n```')["score"], 80)

    def test_model_network_failure_and_empty_response_are_errors(self):
        with patch("app.ollama_client.requests.post", side_effect=requests.Timeout):
            with self.assertRaises(ModelTimeout):
                ask_ollama_chat("system", "resume")

        response = Mock()
        response.status_code = 200
        response.json.return_value = {"message": {"content": ""}}

        with patch("app.ollama_client.requests.post", return_value=response):
            with self.assertRaises(ModelResponseError):
                ask_ollama_chat("system", "resume")


    def test_invalid_metadata_does_not_break_valid_model_response(self):
        response = Mock()
        response.status_code = 200
        response.json.return_value = {
            "message": {"content": '{"score": 80}'},
            "prompt_eval_count": "unknown",
            "eval_count": -1,
            "prompt_eval_duration": "invalid",
            "eval_duration": None,
        }

        with patch("app.ollama_client.requests.post", return_value=response):
            result = ask_ollama_chat("system", "resume")

        self.assertEqual(result["response"], '{"score": 80}')
        self.assertIsNone(result["prompt_tokens"])
        self.assertIsNone(result["generated_tokens"])
        self.assertIsInstance(result["total_time_sec"], float)

    def test_valid_metadata_is_preserved(self):
        response = Mock()
        response.status_code = 200
        response.json.return_value = {
            "message": {"content": '{"score": 80}'},
            "prompt_eval_count": 100,
            "eval_count": 20,
            "prompt_eval_duration": 2_000_000,
            "eval_duration": 500_000,
        }

        with patch("app.ollama_client.requests.post", return_value=response):
            result = ask_ollama_chat("system", "resume")

        self.assertEqual(result["prompt_tokens"], 100)
        self.assertEqual(result["generated_tokens"], 20)


    def test_diagnostic_logs_do_not_contain_full_prompts_or_response(self):
        system_prompt = "SECRET_SYSTEM_PROMPT_123"
        user_prompt = "SECRET_RESUME_TEXT_456"
        model_response = '{"score": 80, "explanation": "SECRET_RESPONSE_789"}'

        response = Mock()
        response.status_code = 200
        response.json.return_value = {
            "message": {"content": model_response},
            "prompt_eval_count": 100,
            "eval_count": 20,
        }

        with patch(
            "app.ollama_client.requests.post",
            return_value=response,
        ), patch("app.ollama_client.log_event") as log:
            result = ask_ollama_chat(system_prompt, user_prompt)

        self.assertEqual(result["response"], model_response)

        diagnostic_calls = [
            call for call in log.call_args_list
            if len(call.args) >= 3
            and call.args[2] in (
                "ollama_request_full",
                "ollama_response_full",
            )
        ]

        self.assertEqual(len(diagnostic_calls), 2)

        for call in diagnostic_calls:
            self.assertNotIn("system_prompt", call.kwargs)
            self.assertNotIn("user_prompt", call.kwargs)
            self.assertNotIn("raw_response", call.kwargs)
            self.assertNotIn("raw_payload", call.kwargs)

    def test_invalid_model_response_is_not_zero_rating(self):
        with patch("app.scoring.ask_ollama_chat", return_value={"response": "bad json"}):
            with self.assertRaises(ModelResponseError):
                rank_profile("resume", self.profile)

    def test_stop_factors_preserve_model_rating_and_cluster_requirements(self):
        rating = {"score": 80, "explanation": "Stop", "triggered_stop_factors": ["stop"]}
        with patch("app.scoring.ask_ollama_chat", return_value={"response": json.dumps(rating)}):
            self.assertEqual(rank_profile("resume", self.profile)["score"], 80)
        self.assertIn("Общие требования", build_ranking_prompt("resume", self.profile, "Общие требования")[1])


    def test_late_valid_model_response_raises_timeout(self):
        rating = {
            "score": 80,
            "explanation": "Valid response",
        }

        def late_response(**kwargs):
            time.sleep(0.05)
            return {"response": json.dumps(rating)}

        with patch("app.scoring.DEFAULT_TIMEOUT", 0.01), \
             patch("app.scoring.ask_ollama_chat", side_effect=late_response):
            with self.assertRaises(ModelTimeout):
                rank_profile(
                    "resume",
                    self.profile,
                    timeout=0.01,
                )

    def test_failure_on_second_profile_does_not_return_partial_result(self):
        requirements = json.dumps({
            "clusters": [{
                "id": "cluster",
                "name": "Test cluster",
                "profiles": [
                    self.profile.model_dump(),
                    {
                        **self.profile.model_dump(),
                        "id": "profile-2",
                        "position_name": "Data Engineer",
                    },
                ],
            }],
        })

        first_result = {
            "score": 80,
            "explanation": "First profile evaluated",
        }

        with patch(
            "app.main.rank_profile",
            side_effect=[
                first_result,
                ModelResponseError("Invalid response for second profile"),
            ],
        ) as rank:
            with self.assertRaises(HTTPException) as context:
                rank_resume(self.upload(), requirements)

        self.assertEqual(context.exception.status_code, 502)
        self.assertEqual(rank.call_count, 2)
        self.assertFalse(model_lock.locked())
    
    def assert_rank_error(self, error, expected_status):
        with patch("app.main.rank_profile", side_effect=error):
            with self.assertRaises(HTTPException) as context:
                rank_resume(self.upload(), self.requirements)

        self.assertEqual(context.exception.status_code, expected_status)
        self.assertFalse(model_lock.locked())

    def test_model_unavailable_returns_503_and_releases_lock(self):
        self.assert_rank_error(
            ModelUnavailable("Ollama unavailable"),
            503,
        )

    def test_model_timeout_returns_504_and_releases_lock(self):
        self.assert_rank_error(
            ModelTimeout("Request timed out"),
            504,
        )

    def test_invalid_model_response_returns_502_and_releases_lock(self):
        self.assert_rank_error(
            ModelResponseError("Invalid model response"),
            502,
        )

    def test_context_fit_error_returns_422_and_releases_lock(self):
        self.assert_rank_error(
            ContextFitError("Context limit exceeded"),
            422,
        )


    def test_success_preserves_contract_and_details(self):
        rating = {"score": 80, "explanation": "OK", "details": {"method": "ollama-agent-v1"}}
        with patch("app.main.rank_profile", return_value=rating) as rank:
            result = rank_resume(self.upload(), self.requirements)
        self.assertEqual(result.clusters[0].profiles[0].score, 80)
        self.assertEqual(result.clusters[0].profiles[0].details["method"], "ollama-agent-v1")
        self.assertEqual(rank.call_args.kwargs["cluster_requirements"], "Общие требования")

    def test_profiles_are_ranked_sequentially_with_cluster_requirements(self):
        req = json.loads(self.requirements)
        req["clusters"][0]["profiles"].append({
            **self.profile.model_dump(),
            "id": "second",
            "position_name": "Data Engineer",
        })
        requirements = json.dumps(req)
        calls = []

        def fake_rank(resume_text, profile, cluster_requirements=None, **kwargs):
            calls.append((profile.id, cluster_requirements))
            return {
                "score": 80,
                "explanation": "OK",
            }

        with patch("app.main.rank_profile", side_effect=fake_rank):
            result = rank_resume(self.upload(), requirements)

        self.assertEqual(len(result.clusters[0].profiles), 2)
        self.assertEqual([item[0] for item in calls], [
            self.profile.id,
            "second",
        ])
        self.assertTrue(all(
            item[1] == req["clusters"][0]["comment"]
            for item in calls
        ))

    def test_partial_failure_returns_no_success_and_disables_analysis(self):
        req = json.loads(self.requirements)
        req["clusters"][0]["profiles"].append({**self.profile.model_dump(), "id": "second"})
        with patch("app.main.rank_profile", side_effect=[{"score": 80, "explanation": "OK"}, ModelUnavailable("offline")]):
            with self.assertRaises(HTTPException) as exc:
                rank_resume(self.upload(), json.dumps(req))
        self.assertEqual(exc.exception.status_code, 503)
        self.assertEqual(readiness().status_code, 503)
        self.assertEqual(health_check(), {"ok": True})

    def test_offline_and_recovery(self):
        mark_ready(False)
        with self.assertRaises(HTTPException) as exc:
            rank_resume(self.upload(), self.requirements)
        self.assertEqual(exc.exception.status_code, 503)
        self.assertEqual(health_check(), {"ok": True})
        mark_ready(True)
        self.assertEqual(readiness().status_code, 200)

    def test_readiness_matches_backend_contract_and_liveness_is_independent(self):
        for available in (False, True):
            mark_ready(available)
            response = readiness()
            self.assertEqual(response.status_code, 200 if available else 503)
            self.assertEqual(json.loads(response.body), {"available": available})
            self.assertEqual(health_check(), {"ok": True})

    
    
    def test_monitor_checks_model_tags_without_generation(self):
        stop = threading.Event()
        observed = []

        responses = [
            requests.ConnectionError("offline"),
            Mock(
                status_code=200,
                json=Mock(
                    return_value={
                        "models": [{"name": "qwen3.5:9b"}]
                    }
                ),
            ),
        ]

        def get_response(*args, **kwargs):
            reply = responses.pop(0)
            if isinstance(reply, Exception):
                raise reply
            return reply

        def wait(_interval):
            observed.append(
                json.loads(readiness().body)["available"]
            )
            if not responses:
                stop.set()

        with patch(
            "app.readiness.requests.get",
            side_effect=get_response,
        ) as get, patch(
            "app.readiness.requests.post"
        ) as post, patch.object(
            stop, "wait", side_effect=wait
        ):
            monitor(stop)

        self.assertEqual(observed, [False, True])
        self.assertEqual(get.call_count, 2)
        post.assert_not_called()
        self.assertFalse(model_lock.locked())

    def test_monitor_skips_probe_while_ranking(self):
        stop = threading.Event()
        model_lock.acquire()
        try:
            with patch("app.readiness.requests.get") as get, \
                 patch.object(stop, "wait", side_effect=lambda _interval: stop.set()):
                monitor(stop)
            get.assert_not_called()
            self.assertEqual(readiness().status_code, 200)
        finally:
            model_lock.release()

    def test_malformed_chat_responses_disable_analysis(self):
        for reply in (None, [], {}, {"message": None}, {"message": {"content": " "}}):
            response = Mock(status_code=200)
            response.json.return_value = reply
            with self.subTest(reply=reply), \
                 patch("app.ollama_client.requests.post", return_value=response):
                with self.assertRaises(ModelResponseError):
                    ask_ollama_chat("system", "resume")
        response = Mock(status_code=200)
        response.json.side_effect = ValueError("invalid JSON")
        with patch("app.ollama_client.requests.post", return_value=response):
            with self.assertRaises(ModelResponseError):
                ask_ollama_chat("system", "resume")

    def test_input_errors_do_not_mark_model_unavailable(self):
        with self.assertRaises(HTTPException) as exc:
            rank_resume(self.upload(), "not json")
        self.assertEqual(exc.exception.status_code, 422)
        self.assertEqual(readiness().status_code, 200)

    def test_extracts_docx_table_text(self):
        doc = Document()
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Python SQL"
        buf = BytesIO()
        doc.save(buf)
        self.assertIn("Python SQL", extract_text_from_file(buf.getvalue(), "resume.docx"))


class LifecycleTests(unittest.IsolatedAsyncioTestCase):
    async def test_startup_runs_monitor_and_shutdown_signals_stop(self):
        started = threading.Event()
        stopped = threading.Event()

        def fake_monitor(stop):
            mark_ready(True)
            started.set()
            stop.wait(2)
            if stop.is_set():
                stopped.set()

        try:
            with patch("app.main.monitor", side_effect=fake_monitor):
                async with lifespan(app):
                    self.assertTrue(await asyncio.to_thread(started.wait, 1))
                    self.assertEqual(json.loads(readiness().body), {"available": True})
                self.assertTrue(await asyncio.to_thread(stopped.wait, 1))
        finally:
            mark_ready(False)


if __name__ == "__main__":
    unittest.main()
