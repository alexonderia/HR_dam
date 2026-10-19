"""Извлечение текста из PDF и DOCX с кэшированием."""

import hashlib
import logging
from collections import OrderedDict
from io import BytesIO

import pymupdf as fitz
from docx import Document

from app.config import settings
from app.logging_config import get_logger, log_event

import subprocess
import tempfile
from pathlib import Path

from striprtf.striprtf import rtf_to_text

logger = get_logger("app.text_extractor")

# LRU-кэш: {хэш_файла: извлечённый_текст}
_text_cache: "OrderedDict[str, str]" = OrderedDict()
_text_cache_bytes = 0

def get_file_hash(content: bytes) -> str:
    """Возвращает SHA-256 хэш содержимого файла."""
    return hashlib.sha256(content).hexdigest()


def _cache_get(file_hash: str) -> str | None:
    if file_hash in _text_cache:
        _text_cache.move_to_end(file_hash)
        return _text_cache[file_hash]
    return None



def _cache_put(file_hash: str, text: str) -> None:
    global _text_cache_bytes

    max_items = settings.TEXT_CACHE_MAX_ITEMS
    max_bytes = settings.TEXT_CACHE_MAX_BYTES
    text_size = len(text.encode("utf-8"))

    # Не кэшируем один текст, который сам превышает лимит.
    if max_items <= 0 or max_bytes <= 0 or text_size > max_bytes:
        return

    # Если запись уже существует, освобождаем её прежний размер.
    old_text = _text_cache.pop(file_hash, None)
    if old_text is not None:
        _text_cache_bytes -= len(old_text.encode("utf-8"))

    _text_cache[file_hash] = text
    _text_cache_bytes += text_size

    # Удаляем самые старые записи, пока не соблюдены оба ограничения.
    while _text_cache and (
        len(_text_cache) > max_items
        or _text_cache_bytes > max_bytes
    ):
        _, removed_text = _text_cache.popitem(last=False)
        _text_cache_bytes -= len(removed_text.encode("utf-8"))



def extract_text_from_pdf_bytes(pdf_bytes: bytes) -> str:
    """Извлекает текст из PDF."""
    try:
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        text = ""
        for page in doc:
            text += page.get_text()
        doc.close()
        return text.strip()
    except Exception as e:
        raise Exception(f"Ошибка при извлечении текста из PDF: {str(e)}")


def extract_text_from_docx_bytes(docx_bytes: bytes) -> str:
    """Извлекает текст из DOCX."""
    try:
        doc = Document(BytesIO(docx_bytes))
        text = ""
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                text += " | ".join(cell.text for cell in row.cells) + "\n"
        return text.strip()
    except Exception as e:
        raise Exception(f"Ошибка при извлечении текста из DOCX: {str(e)}")



def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """
    Извлекает текст из PDF, DOCX, DOC, RTF и TXT.

    Повторно использует текст из LRU-кэша,
    если файл с таким содержимым уже обрабатывался.
    """
    file_hash = get_file_hash(file_bytes)

    cached = _cache_get(file_hash)

    if cached is not None:
        log_event(
            logger,
            logging.INFO,
            "text_cache_hit",
            file_hash=file_hash[:12],
            text_length_chars=len(cached),
        )
        return cached

    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        text = extract_text_from_pdf_bytes(file_bytes)

    elif extension == ".docx":
        text = extract_text_from_docx_bytes(file_bytes)

    elif extension == ".doc":
        text = extract_doc(file_bytes)

    elif extension == ".rtf":
        text = extract_rtf(file_bytes)

    elif extension == ".txt":
        text = extract_txt(file_bytes)

    else:
        raise ValueError(
            f"Неподдерживаемый формат документа: {extension}"
        )

    text = text.strip()

    _cache_put(file_hash, text)

    log_event(
        logger,
        logging.INFO,
        "text_cache_miss",
        file_hash=file_hash[:12],
        text_length_chars=len(text),
    )

    return text


def extract_txt(file_bytes: bytes) -> str:
    """Извлекает текст из TXT."""

    if file_bytes.startswith((b"\xff\xfe", b"\xfe\xff")):
        return file_bytes.decode("utf-16")

    if file_bytes.startswith(b"\xef\xbb\xbf"):
        return file_bytes.decode("utf-8-sig")

    for encoding in ("utf-8", "cp1251"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue

    raise ValueError("Не удалось определить кодировку TXT.")


def extract_rtf(file_bytes: bytes) -> str:
    """Извлекает текст из RTF."""

    # RTF использует ASCII-совместимую разметку,
    # а символы могут быть представлены escape-последовательностями.
    raw_text = file_bytes.decode("latin-1")

    return rtf_to_text(raw_text)


def extract_doc(file_bytes: bytes) -> str:
    """Преобразует DOC в DOCX через LibreOffice."""

    from docx import Document
    from io import BytesIO

    with tempfile.TemporaryDirectory() as tmp:
        directory = Path(tmp)

        source = directory / "resume.doc"
        output = directory / "converted"
        output.mkdir()

        source.write_bytes(file_bytes)

        try:
            subprocess.run(
                [
                    "libreoffice",
                    "-env:UserInstallation=file://" + str(directory / "lo-profile"),
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    str(output),
                    str(source),
                ],
                check=True,
                capture_output=True,
                timeout=60,
            )
        except (
            subprocess.CalledProcessError,
            subprocess.TimeoutExpired,
            FileNotFoundError,
        ) as exc:
            raise ValueError(
                "Не удалось преобразовать DOC в DOCX."
            ) from exc

        converted = output / "resume.docx"

        if not converted.exists():
            raise ValueError(
                "LibreOffice не создал DOCX."
            )

        return extract_text_from_docx_bytes(converted.read_bytes())
